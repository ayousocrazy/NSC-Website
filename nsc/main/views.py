from django.shortcuts import render, get_object_or_404, redirect
from django.db.models import Case, When, IntegerField
from .models import *
import os
from openpyxl import Workbook, load_workbook
from django.conf import settings
import portalocker 
from django.contrib import messages
from django.http import HttpResponse
from django.db import connection
from django.db.utils import OperationalError

EXCEL_FILE = os.path.join(settings.BASE_DIR, 'main', 'excel', 'admissions.xlsx')


def home(request):
    programs = Academics.objects.filter(level="bachelors").order_by(
        Case(
            When(program_key="csit", then=0),
            When(program_key="bca", then=1),
            default=2,
            output_field=IntegerField()
        ),
        "title"
    )
    default_program = programs.first()
    faqs = FAQs.objects.filter(level="bachelors").order_by("created")

    context = {
        "plus2": False,
        "programs": programs,
        "default_program": default_program,
        "faqs": faqs
    }
    return render(request, "main/home.html", context)


def homePlus2(request):
    programs = Academics.objects.filter(level="plus2").order_by(
        Case(
            When(program_key="science", then=0),
            When(program_key="management", then=1),
            default=2,
            output_field=IntegerField()
        ),
        "title"
    )
    default_program = programs.first()
    faqs = FAQs.objects.filter(level="plus2").order_by("created")

    context = {
        "plus2": True,
        "programs": programs,
        "default_program": default_program,
        "faqs": faqs
    }
    return render(request, "main/home.html", context)


def academics(request, pk):
    program = get_object_or_404(Academics, program_key=pk)

    if program.level == "plus2":
        core_subjects = SubjectPlus2.objects.filter(program=program, optional=False)
        optional_subjects = SubjectPlus2.objects.filter(program=program, optional=True)
        careers = CareerPlus2.objects.filter(program=program)
        program.criteria = program.criteria.split('\n') if program.criteria else []

        context = {
            "plus2": True,
            "program": program,
            "core_subjects": core_subjects,
            "optional_subjects": optional_subjects,
            "careers": careers,
        }
        return render(request, "main/plus2academics.html", context)

    courses = Course.objects.filter(program=program).order_by('year', 'semester')
    careers = Career.objects.filter(program=program)
    program.criteria = program.criteria.split('\n') if program.criteria else []

    context = {
        "plus2": False,
        "program": program,
        "courses": courses,
        "careers": careers,
    }

    return render(request, "main/academics.html", context)


def admissions(request):
    plus2 = request.GET.get('level') == 'plus2'
    return render(request, "main/admissions.html", {"plus2": plus2})


def faculty(request):
    plus2 = request.GET.get('level') == 'plus2'
    management_faculties = Faculty.objects.filter(
        category__in=['MANAGEMENT', 'HOD']
    )
    teaching_faculties = Faculty.objects.filter(category='TEACHING')
    non_teaching_faculties = Faculty.objects.filter(category='NON_TEACHING')

    context = {
        'plus2': plus2,
        'management_faculties': management_faculties,
        'teaching_faculties': teaching_faculties,
        'non_teaching_faculties': non_teaching_faculties,
    }
    return render(request, "main/faculty.html", context)

def about(request):
    plus2 = request.GET.get('level') == 'plus2'
    return render(request, "main/about.html", {'plus2': plus2})

def form(request):
    if request.method == "POST":
        try:
            data = [
                request.POST.get("first_name"),
                request.POST.get("middle_name"),
                request.POST.get("last_name"),
                request.POST.get("gender"),
                request.POST.get("email"),
                request.POST.get("phone"),
                request.POST.get("address"),
                request.POST.get("program"),
                request.POST.get("sub_stream"),
                request.POST.get("query"),
            ]

            excel_dir = os.path.dirname(EXCEL_FILE)
            os.makedirs(excel_dir, exist_ok=True)

            if not os.path.exists(EXCEL_FILE):
                wb = Workbook()
                ws = wb.active
                ws.title = "Admissions"
                ws.append([
                    "First Name", "Middle Name", "Last Name", "Gender",
                    "Email", "Phone", "Address", "Program",
                    "Sub Stream", "Query"
                ])
                wb.save(EXCEL_FILE)

            with open(EXCEL_FILE, 'rb+') as f:
                portalocker.lock(f, portalocker.LOCK_EX)
                try:
                    wb = load_workbook(f)
                except Exception:
                    wb = Workbook()
                    ws = wb.active
                    ws.title = "Admissions"
                    ws.append([
                        "First Name", "Middle Name", "Last Name", "Gender",
                        "Email", "Phone", "Address", "Program",
                        "Sub Stream", "Query"
                    ])
                ws = wb.active
                ws.append(data)
                f.seek(0)
                wb.save(f)
                portalocker.unlock(f)

            messages.success(request, "Form submitted successfully. Our team will contact you.")
            program = request.POST.get("program")
            if program == "Bachelors":
                return redirect("/")
            elif program == "Plus2":
                return redirect("/plus2/")
            else:
                return redirect("/")

        except Exception as e:
            print("FORM ERROR:", e)
            messages.error(request, "Submission failed. Please try again.")
            return redirect(request.path)  

    return render(request, "main/form.html", {"no_footer": True})

def downloadAdmissionsList(request):
    try:
        with open(EXCEL_FILE, "rb") as file:
            response = HttpResponse(
                file.read(),
                content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            response["Content-Disposition"] = 'attachment; filename="AdmissionList.xlsx"'
            return response
    except FileNotFoundError:
        messages.error(request, "Admission file not found.")
        return redirect("/")
    except Exception as e:
        print("DOWNLOAD ERROR:", e)
        messages.error(request, "Could not download file. Please try again.")
        return redirect("/")

# def test_db(request):
#     try:
#         with connection.cursor() as cursor:
#             cursor.execute("SELECT 1;")
#             result = cursor.fetchone()
#         return HttpResponse(f"Database connection OK! Result: {result}")
#     except OperationalError as e:
#         return HttpResponse(f"Database connection FAILED! Error: {e}")